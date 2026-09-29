"""Research public executive contacts for the Canadian WP6 workbook.

Only direct official team-card evidence can be accepted automatically. Other
candidates require review before they reach the exported workbook.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import time
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import quote, urljoin, urlsplit
from urllib.request import Request, urlopen

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from bbt_bizdev.adapters.linkedin import canonicalize_linkedin_url, configured_search

SOURCE = Path("outputs/canada_company_enrichment_wp6_2026-07-30/Canada_Company_Enrichment_Work_Package_6.xlsx")
OUT = Path("outputs/canada_executive_research")
FIELDS = ["candidate_id", "row_id", "company", "name", "title", "linkedin_url", "email", "phone", "role_url", "linkedin_evidence_url", "email_evidence_url", "phone_evidence_url", "identity_verified", "employment_verified", "linkedin_verified", "decision", "reason", "checked_at", "evidence_record_id"]
ROLE_PATTERNS = [("CEO", r"\b(?:ceo|chief executive officer|chef de la direction)\b"), ("COO", r"\b(?:coo|chief operating officer|chef des opérations)\b"), ("CTO", r"\b(?:cto|chief technology officer|chief technical officer|chef de la technologie)\b"), ("CFO", r"\b(?:cfo|chief financial officer|chef des finances)\b"), ("Other", r"\b(?:president|président|managing director|chief [a-z ]+ officer|vice president|vice-président|vp)\b")]
NAV_WORDS = {"about", "contact", "team", "leadership", "people", "company", "equipe", "notre equipe", "a propos", "direction"}


class PageParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.raw = ""
        self.links = []
        self.text = []
        self.href = None
        self.anchor = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.anchor = []

    def handle_data(self, data):
        clean = " ".join(data.split())
        if clean:
            self.text.append(clean)
            if self.href:
                self.anchor.append(clean)

    def handle_endtag(self, tag):
        if tag == "a" and self.href:
            self.links.append((self.href, " ".join(self.anchor)))
            self.href = None


def read_rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    companies = []
    for number, row in enumerate(book["Companies"].iter_rows(min_row=2, values_only=True), 2):
        if not row[1]:
            continue
        listed_website = str(row[2] or "").strip()
        companies.append({"row_id": f"wp6-row-{number}", "source_row": number, "rank": row[0], "company": str(row[1]).strip(), "website": listed_website, "source_website": listed_website, "province": row[4] or "", "product": row[6] or "", "identity_status": "unverified", "status": "not_started"})
    reviewed = set()
    unresolved = set()
    alias_records = {}
    for alias_row, row in enumerate(book["Alias Review"].iter_rows(min_row=2, values_only=True), 2):
        names = [str(row[3] or "").strip()]
        if row[1]:
            names.extend(part.strip() for part in str(row[1]).split(";"))
        if row[0] == "Ambiguous name" and row[4]:
            try:
                names.extend(json.loads(row[4]))
            except (ValueError, TypeError):
                pass
        names = {name.casefold() for name in names if name}
        reviewed.update(names)
        if row[0] == "Ambiguous name":
            unresolved.update(names)
            try:
                detail = json.loads(row[7]) if row[7] else {}
                review = {"source_sheet": "Alias Review", "source_row": alias_row,
                          "issue_type": detail.get("issue_type", ""),
                          "names": json.loads(row[4] or "[]"),
                          "domains": json.loads(row[5] or "[]")}
            except (ValueError, TypeError):
                review = {"issue_type": "unparsed", "names": [], "domains": []}
            for name in names:
                alias_records.setdefault(name, []).append(review)
    for company in companies:
        key = company["company"].casefold()
        company["alias_review"] = key in reviewed
        company["identity_issue"] = key in unresolved
        company["identity_reviews"] = alias_records.get(key, [])
    return companies


def sync_identity_metadata(companies, source_path):
    source_by_id = {row["row_id"]: row for row in read_rows(source_path)}
    for company in companies:
        source = source_by_id.get(company["row_id"])
        if source:
            company["alias_review"] = source["alias_review"]
            company["identity_issue"] = source["identity_issue"]
            company["identity_reviews"] = source["identity_reviews"]
            company["source_website"] = source["source_website"]
            if company.get("website_source") == "luna_discovery_verified_homepage":
                company["identity_status"] = "homepage_name_match"
                company["errors"] = [item for item in company.get("errors", []) if item != "missing_website"]


def pilot(rows):
    selected = []
    seen = set()
    def take(items, n):
        if n <= 0:
            return
        count = 0
        for row in items:
            if row["row_id"] not in seen:
                selected.append(row); seen.add(row["row_id"]); count += 1
                if count == n:
                    break
    take(rows, 10)
    take((r for r in rows if not r.get("source_website", r["website"])), 5)
    take((r for r in rows if r.get("alias_review", r["identity_issue"])), 5)
    stride = max(1, len(rows) // 5)
    take((rows[i] for i in range(0, len(rows), stride)), 5)
    take(rows, 25 - len(selected))
    return selected


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


def load_json(path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def write_csv(path, rows, fields=FIELDS):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fields, extrasaction="ignore")
        writer.writeheader(); writer.writerows(rows)


def read_csv(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def defer_new_automatic_acceptances(rows, accepted_ids, reason):
    """Preserve baseline accepts; route new legacy auto-decisions through investigation."""
    deferred = 0
    for row in rows:
        if row.get("decision") != "accept" or row.get("candidate_id") in accepted_ids:
            continue
        row.update({"decision": "pending", "reason": reason, "identity_verified": "",
                    "employment_verified": "", "linkedin_verified": "", "email": "",
                    "email_evidence_url": "", "phone": "", "phone_evidence_url": ""})
        deferred += 1
    return deferred


def fetch_page(url):
    try:
        request = Request(url, headers={"User-Agent": "Mozilla/5.0 (compatible; BBTResearch/1.0)"})
        with urlopen(request, timeout=15) as response:
            if "html" not in response.headers.get("Content-Type", "").lower():
                return None, "non_html"
            final_url = response.geturl()
            raw = response.read(1_000_000).decode("utf-8", "ignore")
        parser = PageParser(); parser.feed(raw)
        parser.raw = raw
        parser.final_url = final_url
        return parser, ""
    except Exception as error:
        return None, f"{type(error).__name__}: {error}"


def website_pages(website):
    if not website:
        return [], "missing_website"
    root = urlsplit(website)
    if root.scheme not in {"http", "https"} or not root.netloc:
        return [], "invalid_website"
    page, error = fetch_page(website)
    if not page:
        return [], error
    final_url = getattr(page, "final_url", website)
    final_host = urlsplit(final_url).hostname or ""
    links = {}
    for href, label in page.links:
        url = urljoin(final_url, href)
        if urlsplit(url).fragment:
            continue
        if (urlsplit(url).hostname or "").removeprefix("www.").lower() != final_host.removeprefix("www.").lower():
            continue
        clue = f"{label} {urlsplit(url).path}".casefold()
        if not any(word in clue for word in NAV_WORDS):
            continue
        priority = 3 if any(word in clue for word in ("team", "leadership", "people", "equipe")) else (2 if any(word in clue for word in ("about", "a propos")) else 1)
        links[url] = max(priority, links.get(url, 0))
    ordered = sorted(links, key=lambda url: -links[url])
    return [(final_url, page)] + [(url, None) for url in ordered[:6]], ""


def candidate_id(row_id, url, name):
    digest = hashlib.sha256(f"{row_id}|{url}|{name.casefold()}".encode()).hexdigest()[:16]
    return f"exec-{digest}"


def site_candidates(company, url, page):
    # An official LinkedIn link near a name is useful discovery evidence.
    # The reviewer must still verify the person and current role.
    found = []
    for href, anchor in page.links:
        profile = canonicalize_linkedin_url(urljoin(url, href), "person")
        if not profile:
            continue
        name = " ".join(anchor.split())
        if not re.fullmatch(r"[\wÀ-ÿ'.-]+(?:\s+[\wÀ-ÿ'.-]+){1,4}", name):
            name = ""
        found.append({"name": name, "title": "", "linkedin_url": profile, "role_url": url, "linkedin_evidence_url": url, "discovery": "official_link"})
    return found


def clean_words(value):
    return re.sub(r"[^\w]+", " ", value.casefold(), flags=re.UNICODE).strip()


def clean_person_name(value):
    value = value.split(",", 1)[0].strip()
    value = re.sub(r"^(?:Dr\.?\s+)", "", value, flags=re.I)
    value = re.sub(r"(?:[.,]?\s+(?:MD|PhD|MSc|MBA|CPA|CA|FACC|FHRS|FRCSC))+$", "", value, flags=re.I)
    return value.strip(" .,")


def homepage_matches_company(company, page):
    """A workbook website is only trusted when its homepage names the entity."""
    if company.get("identity_issue"):
        return False
    name = clean_words(company["company"])
    if len(name) < 5:
        return False
    soup = BeautifulSoup(page.raw, "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    head = " ".join(str(tag.get_text(" ", strip=True)) for tag in soup.find_all(["title", "h1", "h2"], limit=15))
    body = soup.get_text(" ", strip=True)[:2500]
    return name in clean_words(head + " " + body)


def verified_site_candidates(company, url, page, identity_verified):
    """Accept only a named executive whose personal link is in the same official team card."""
    if not identity_verified:
        return []
    soup = BeautifulSoup(page.raw, "html.parser")
    found = []
    seen = set()
    for anchor in soup.find_all("a", href=True):
        profile = canonicalize_linkedin_url(urljoin(url, anchor["href"]), "person")
        if not profile or profile in seen:
            continue
        for card in list(anchor.parents)[:5]:
            if card.name not in {"article", "li", "div"}:
                continue
            card_text = " ".join(card.stripped_strings)
            if len(card_text) > 600 or len(card_text) < 8:
                continue
            profiles = {
                canonicalize_linkedin_url(urljoin(url, a["href"]), "person")
                for a in card.find_all("a", href=True)
            }
            profiles.discard("")
            if profiles != {profile}:
                continue
            headings = card.find_all(["h2", "h3", "h4", "h5", "strong", "p"], limit=12)
            possible_names = [" ".join(h.stripped_strings) for h in headings
                              if len(" ".join(h.stripped_strings)) <= 90]
            possible_names.insert(0, " ".join(anchor.stripped_strings))
            names = [(raw, clean_person_name(raw)) for raw in possible_names]
            names = [(raw, name) for raw, name in names if role_rank(name) == 99
                     and re.fullmatch(r"[A-ZÀ-Ý][\wÀ-ÿ'.-]+(?:\s+[A-ZÀ-Ý][\wÀ-ÿ'.-]+){1,3}", name)]
            if len({name for _, name in names}) != 1:
                continue
            raw_name, name = max(names, key=lambda pair: len(pair[0]))
            remaining = card_text.replace(raw_name, "", 1)
            if not remaining or role_rank(remaining) == 99:
                continue
            role_tags = [tag.get_text(" ", strip=True) for tag in card.find_all(["h2", "h3", "h4", "h5", "p"], limit=10)]
            role_tags = [value for value in role_tags if value and len(value) <= 100 and name not in value and role_rank(value) != 99]
            title = min(role_tags, key=len) if role_tags else remaining
            title = re.sub(r"^(?:[\s,./]*(?:MD|PhD|MSc|MBA|CPA|CA|FACC|FHRS|FRCSC)\b)+", "", title, flags=re.I)
            title = re.split(r"\bConnect On\b|\bLinkedIn\b|", title, maxsplit=1, flags=re.I)[0].strip(" ,./")
            if not title or len(title) > 100 or role_rank(title) == 99:
                continue
            addresses = {
                a["href"][7:].split("?", 1)[0].strip().casefold()
                for a in card.find_all("a", href=True)
                if a["href"].lower().startswith("mailto:")
            }
            addresses = {address for address in addresses if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", address)
                         and address.split("@", 1)[0] not in {"info", "contact", "sales", "hello", "support", "careers", "admin"}
                         and clean_words(name.split()[-1]) in clean_words(address.split("@", 1)[0])}
            email = next(iter(addresses)) if len(addresses) == 1 else ""
            found.append({"name": name, "title": title, "linkedin_url": profile,
                          "role_url": url, "linkedin_evidence_url": url,
                          "email": email, "email_evidence_url": url if email else "",
                          "discovery": "automated_official_team_card", "auto_verified": True})
            seen.add(profile)
            break
    return found


def search_candidates(company):
    results = []
    queries = [f'"{company}" CEO', f'site:linkedin.com/in "{company}" CEO OR COO OR CTO OR CFO']
    errors = []
    for query in queries:
        hits, error = configured_search(query)
        if error:
            errors.append(f"{query}: {error}")
            if "no parseable results" in error.lower():
                break
        for hit in hits[:5]:
            profile = canonicalize_linkedin_url(hit.url, "person")
            if profile:
                results.append({"name": "", "title": hit.title[:120], "linkedin_url": profile, "role_url": "", "linkedin_evidence_url": hit.url, "discovery": "search_hit", "snippet": hit.snippet[:300]})
    return results, errors


def collect_one(company, use_search=True):
    candidates = []
    pages, error = website_pages(company["website"])
    errors = [error] if error else []
    identity_verified = bool(pages and homepage_matches_company(company, pages[0][1]))
    if pages:
        company["identity_status"] = "homepage_name_match" if identity_verified else "website_needs_review"
    queue = list(pages)
    visited = []
    seen_pages = {url for url, _ in queue}
    for url, page in queue:
        if page is None:
            time.sleep(1)
            page, issue = fetch_page(url)
            if issue:
                errors.append(f"{url}: {issue}")
        if page:
            visited.append(url)
            candidates.extend(site_candidates(company, url, page))
            candidates.extend(verified_site_candidates(company, url, page, identity_verified))
            if len(queue) < 12:
                for href, label in page.links:
                    nested = urljoin(getattr(page, "final_url", url), href)
                    parts = urlsplit(nested)
                    clue = f"{label} {parts.path}".casefold()
                    if (parts.fragment or nested in seen_pages
                            or not any(word in clue for word in ("team", "leadership", "people", "equipe"))
                            or (parts.hostname or "").removeprefix("www.").lower()
                            != (urlsplit(url).hostname or "").removeprefix("www.").lower()):
                        continue
                    queue.append((nested, None))
                    seen_pages.add(nested)
                    if len(queue) >= 12:
                        break
    if use_search and sum(item.get("auto_verified", False) for item in candidates) < 3:
        searched, search_errors = search_candidates(company["company"])
        candidates.extend(searched); errors.extend(search_errors)
    unique = {}
    for item in candidates:
        profile = item["linkedin_url"]
        if profile not in unique or item.get("auto_verified") or (item["discovery"] == "official_link" and unique[profile]["discovery"] == "search_hit"):
            unique[profile] = item
    output = []
    for item in unique.values():
        automatic = item.get("auto_verified", False)
        output.append({"candidate_id": candidate_id(company["row_id"], item["linkedin_url"], item["name"]), "row_id": company["row_id"], "company": company["company"], "name": item["name"], "title": item["title"], "linkedin_url": item["linkedin_url"], "email": item.get("email", ""), "phone": "", "role_url": item["role_url"], "linkedin_evidence_url": item["linkedin_evidence_url"], "email_evidence_url": item.get("email_evidence_url", ""), "phone_evidence_url": "", "identity_verified": "yes" if automatic else "", "employment_verified": "yes" if automatic else "", "linkedin_verified": "yes" if automatic else "", "decision": "accept" if automatic else "pending", "reason": item["discovery"], "checked_at": date.today().isoformat()})
    company["status"] = "auto_verified" if any(item["decision"] == "accept" for item in output) else ("needs_review" if output else ("no_candidates" if pages else ("website_missing" if error == "missing_website" else "site_unavailable")))
    company["errors"] = errors
    company["pages_visited"] = visited
    return output


def role_rank(title):
    for index, (_, pattern) in enumerate(ROLE_PATTERNS):
        if re.search(pattern, title or "", re.I):
            return index
    return 99


def validate(rows, decisions):
    by_id = {r["row_id"]: r for r in rows}
    valid = []
    issues = []
    for item in decisions:
        if item.get("decision") != "accept":
            continue
        required = ["row_id", "name", "title", "linkedin_url", "role_url", "linkedin_evidence_url"]
        missing = [key for key in required if not item.get(key)]
        for key in ("role_url", "linkedin_evidence_url", "email_evidence_url", "phone_evidence_url"):
            value = item.get(key, "")
            if value and (urlsplit(value).scheme not in {"http", "https"} or not urlsplit(value).netloc):
                missing.append(f"valid {key}")
        missing.extend(key for key in ("identity_verified", "employment_verified", "linkedin_verified") if item.get(key, "").casefold() != "yes")
        if item.get("row_id") not in by_id:
            missing.append("known company row")
        if role_rank(item.get("title", "")) == 99:
            missing.append("eligible current executive title")
        if not canonicalize_linkedin_url(item.get("linkedin_url", ""), "person"):
            missing.append("valid LinkedIn person URL")
        if item.get("email") and not item.get("email_evidence_url"):
            missing.append("email evidence URL")
        if item.get("email"):
            address = item["email"].strip().casefold()
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", address) or address.split("@", 1)[0] in {"info", "contact", "sales", "hello", "support", "careers", "admin"}:
                missing.append("individual business email")
        if item.get("phone") and not item.get("phone_evidence_url"):
            missing.append("phone evidence URL")
        if missing:
            issues.append(f"{item.get('candidate_id')}: missing/invalid {', '.join(missing)}")
        else:
            valid.append(item)
    seen = set()
    selected = []
    for item in sorted(valid, key=lambda x: (int(x["row_id"].split("-")[-1]), role_rank(x["title"]), x["name"])):
        key = (item["row_id"], canonicalize_linkedin_url(item["linkedin_url"], "person"))
        if key in seen:
            issues.append(f"{item['candidate_id']}: duplicate person URL")
            continue
        seen.add(key)
        if sum(x["row_id"] == item["row_id"] for x in selected) < 3:
            selected.append(item)
    return selected, issues


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["prepare", "collect", "check-openrouter", "luna", "investigate", "review", "validate"])
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--pilot", action="store_true")
    parser.add_argument("--blind-replay", action="store_true",
                        help="Investigate three hidden known-good contacts and five pilot traps")
    parser.add_argument("--limit", type=int, default=50)
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--max-cost-usd", type=float, default=1.0,
                        help="Luna pass: stop launching calls after this reported spend (default: 1.0)")
    parser.add_argument("--max-requests", type=int, default=24)
    parser.add_argument("--max-searches", type=int, default=48)
    parser.add_argument("--max-fetches", type=int, default=32)
    parser.add_argument("--rows", type=str, default="",
                        help="Comma-separated source worksheet row numbers for a targeted collect or Luna run")
    parser.add_argument("--no-search", action="store_true",
                        help="Collect from official sites only when the public search service is unavailable")
    parser.add_argument("--key-file", type=Path,
                        help="Private OpenRouter key file; defaults to ~/.config/bbt_bizdev/openrouter.key")
    args = parser.parse_args()
    try:
        chosen_rows = {int(part.strip()) for part in args.rows.split(",") if part.strip()}
    except ValueError:
        parser.error("--rows must contain comma-separated source worksheet row numbers")
    path = args.out
    if args.command == "check-openrouter":
        from bbt_bizdev.canada_executive_luna import check_openrouter_key
        try:
            print(json.dumps(check_openrouter_key(args.key_file)))
        except RuntimeError as error:
            parser.error(str(error))
    elif args.command == "prepare":
        if (path / "companies.json").exists() and not args.refresh:
            parser.error("Research state exists. Use --refresh to replace it, or continue with collect/review/validate")
        companies = read_rows(args.source)
        save_json(path / "companies.json", companies)
        save_json(path / "manifest.json", {"source": str(args.source), "sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(), "prepared_at": date.today().isoformat(), "company_count": len(companies)})
        write_csv(path / "decisions.csv", [])
        print(f"Prepared {len(companies)} companies; pilot has {len(pilot(companies))} rows")
    elif args.command == "collect":
        companies = load_json(path / "companies.json", [])
        if not companies:
            parser.error("Run prepare first")
        sync_identity_metadata(companies, args.source)
        targets = pilot(companies) if args.pilot else companies
        if chosen_rows:
            targets = [company for company in targets if company["source_row"] in chosen_rows]
        candidates = read_csv(path / "candidates.csv")
        baseline_accepted_ids = {r["candidate_id"] for r in read_csv(path / "decisions.csv")
                                 if r.get("decision") == "accept"}
        existing = {r["candidate_id"] for r in candidates}
        processed = 0
        for company in targets:
            if processed >= args.limit:
                break
            if company["status"] != "not_started" and not args.refresh:
                continue
            if args.refresh:
                candidates = [item for item in candidates if item["row_id"] != company["row_id"]]
                existing = {item["candidate_id"] for item in candidates}
            found = collect_one(company, use_search=not args.no_search)
            defer_new_automatic_acceptances(found, baseline_accepted_ids,
                                            "collector_requires_investigation")
            for item in found:
                if item["candidate_id"] not in existing:
                    candidates.append(item); existing.add(item["candidate_id"])
            write_csv(path / "candidates.csv", candidates)
            decisions = read_csv(path / "decisions.csv")
            by_decision_id = {item["candidate_id"]: index for index, item in enumerate(decisions)}
            for item in found:
                if item["decision"] != "accept":
                    if (item.get("reason") == "collector_requires_investigation"
                            and item["candidate_id"] not in by_decision_id):
                        decisions.append(item)
                    continue
                old_index = by_decision_id.get(item["candidate_id"])
                if old_index is None:
                    decisions.append(item)
                elif (decisions[old_index].get("decision") == "pending"
                      or decisions[old_index].get("reason") == "automated_official_team_card"):
                    decisions[old_index] = item
            company_decisions = [item for item in decisions if item["row_id"] == company["row_id"]]
            if any(item["decision"] == "accept" for item in company_decisions):
                company["status"] = "auto_verified"
            elif any(item["decision"] == "pending" for item in company_decisions):
                company["status"] = "needs_review"
            write_csv(path / "decisions.csv", decisions)
            save_json(path / "companies.json", companies)
            processed += 1
            print(f"{company['row_id']} {company['company']}: {company['status']}, {len(found)} candidates", flush=True)
            if any("Name or service not known" in message for message in company.get("errors", [])):
                print("Stopping batch: DNS resolution is unavailable; remaining companies are untouched", flush=True)
                break
        print(f"Processed {processed} companies")
    elif args.command == "luna":
        from bbt_bizdev.canada_executive_luna import run_luna_pass
        companies = load_json(path / "companies.json", [])
        if not companies:
            parser.error("Run prepare first")
        sync_identity_metadata(companies, args.source)
        decisions = read_csv(path / "decisions.csv")
        baseline_accepted_ids = {r["candidate_id"] for r in decisions if r.get("decision") == "accept"}
        try:
            decisions, summary = run_luna_pass(
                companies, decisions, path, pilot_only=args.pilot,
                limit=args.limit, max_cost_usd=args.max_cost_usd,
                key_file=args.key_file, source_rows=chosen_rows,
            )
        except (RuntimeError, ValueError) as error:
            parser.error(str(error))
        deferred = defer_new_automatic_acceptances(decisions, baseline_accepted_ids,
                                                    "legacy_luna_requires_investigation")
        summary["deferred_by_investigation_gate"] = deferred
        summary["accepted"] -= deferred
        write_csv(path / "decisions.csv", decisions)
        save_json(path / "companies.json", companies)
        save_json(path / "luna_last_run.json", summary)
        print(json.dumps(summary))
    elif args.command == "investigate":
        from bbt_bizdev.canada_executive_investigation import (
            BLIND_HIDDEN_NAMES, BLIND_REPLAY_ROWS, PILOT_ROWS, run_investigation,
        )
        if args.blind_replay:
            if path.resolve() == OUT.resolve():
                parser.error("--blind-replay requires an isolated --out directory")
            if chosen_rows and not chosen_rows.issubset(BLIND_REPLAY_ROWS):
                parser.error("--blind-replay --rows must select replay rows")
            chosen_rows = chosen_rows or set(BLIND_REPLAY_ROWS)
        elif args.pilot and not chosen_rows:
            chosen_rows = set(PILOT_ROWS)
        if not chosen_rows:
            parser.error("investigate requires --pilot or explicit --rows")
        if path.resolve() != OUT.resolve() and not (path / "companies.json").exists():
            path.mkdir(parents=True, exist_ok=True)
            for filename in ("companies.json", "decisions.csv", "manifest.json"):
                shutil.copy2(OUT / filename, path / filename)
        if args.blind_replay and not (path / "blinded_truth.json").exists():
            all_decisions = read_csv(path / "decisions.csv")
            hidden = []
            for source_row, name in BLIND_HIDDEN_NAMES:
                matches = [row for row in all_decisions if row.get("row_id") == f"wp6-row-{source_row}"
                           and row.get("name", "").casefold() == name.casefold()
                           and row.get("decision") == "accept"]
                if len(matches) != 1:
                    parser.error(f"Blinded contact not uniquely accepted: row {source_row}, {name}")
                hidden.append({key: matches[0][key] for key in
                               ("candidate_id", "row_id", "name", "title", "linkedin_url")})
            hidden_ids = {row["candidate_id"] for row in hidden}
            write_csv(path / "decisions.csv", [row for row in all_decisions
                                               if row["candidate_id"] not in hidden_ids])
            save_json(path / "blinded_truth.json", {"hidden": hidden,
                "trap_rows": [f"wp6-row-{row}" for row in BLIND_REPLAY_ROWS[3:]]})
        companies = load_json(path / "companies.json", [])
        if not companies:
            parser.error("Run prepare first")
        sync_identity_metadata(companies, args.source)
        decisions = read_csv(path / "decisions.csv")
        try:
            decisions, summary = run_investigation(
                companies, decisions, path, source_rows=chosen_rows,
                max_cost_usd=args.max_cost_usd, max_requests=args.max_requests,
                max_searches=args.max_searches, max_fetches=args.max_fetches,
                key_file=args.key_file,
                target_order=BLIND_REPLAY_ROWS if args.blind_replay else
                PILOT_ROWS if args.pilot else None,
            )
        except (RuntimeError, ValueError) as error:
            parser.error(str(error))
        write_csv(path / "decisions.csv", decisions)
        save_json(path / "companies.json", companies)
        selected, issues = validate(companies, decisions)
        save_json(path / "selected.json", selected)
        summary["selected_contacts"] = len(selected)
        summary["newly_selected_vs_37"] = len(selected) - 37
        if args.blind_replay:
            truth = load_json(path / "blinded_truth.json", {})
            summary["hidden_known_good"] = len(truth["hidden"])
            summary["recovered_known_good"] = sum(any(
                row["row_id"] == hidden["row_id"] and row["name"].casefold() == hidden["name"].casefold()
                and canonicalize_linkedin_url(row["linkedin_url"], "person") ==
                canonicalize_linkedin_url(hidden["linkedin_url"], "person")
                for row in selected) for hidden in truth["hidden"])
            summary["false_accepts_in_trap_rows"] = [
                {"row_id": row["row_id"], "name": row["name"]} for row in decisions
                if row.get("decision") == "accept" and row.get("row_id") in truth["trap_rows"]]
        summary["validation_issues"] = issues
        save_json(path / "investigation_last_run.json", summary)
        print(json.dumps(summary))
    elif args.command == "review":
        old = {r["candidate_id"]: r for r in read_csv(path / "decisions.csv")}
        candidates = read_csv(path / "candidates.csv")
        current_auto_ids = {r["candidate_id"] for r in candidates if r.get("reason") == "automated_official_team_card"}
        updated = []
        for candidate in candidates:
            prior = old.get(candidate["candidate_id"])
            updated.append(candidate if candidate.get("reason") == "automated_official_team_card" and (not prior or prior.get("reason") == "automated_official_team_card") else (prior or candidate))
        candidate_ids = {r["candidate_id"] for r in updated}
        updated.extend(row for key, row in old.items() if key not in candidate_ids and (
            row.get("reason") != "automated_official_team_card" or key in current_auto_ids
        ))
        updated = [row for row in updated if row.get("reason") != "automated_official_team_card" or row["candidate_id"] in current_auto_ids]
        write_csv(path / "decisions.csv", updated)
        print(f"Review queue has {len(updated)} candidates")
    else:
        companies = load_json(path / "companies.json", [])
        selected, issues = validate(companies, read_csv(path / "decisions.csv"))
        save_json(path / "selected.json", selected)
        print(f"Selected {len(selected)} contacts; {len(issues)} issues")
        for issue in issues:
            print(issue)
        if issues:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
