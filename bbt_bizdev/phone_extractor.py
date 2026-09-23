from __future__ import annotations

import csv
import hashlib
import json
import re
import ssl
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, field
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable
from http.client import IncompleteRead
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlsplit, urlunsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from .adapters.linkedin import PublicSearchHit, configured_search, duckduckgo_search
from .config import USER_AGENT


PHONE_RE = re.compile(
    r"(?<![\w\d])(?:\+\s*353|00\s*353|0)(?:[\s()./\-]*\d){7,11}(?!\d)",
    re.I,
)
PHONE_CONTEXT_RE = re.compile(
    r"\b(?:tel(?:ephone)?|phone|mobile|call|contact|office|reception|switchboard|sales|support|enquiries)\b",
    re.I,
)
FAX_CONTEXT_RE = re.compile(r"\bfax\b", re.I)
THIRD_PARTY_CONTEXT_RE = re.compile(
    r"(?:dataprotection\.ie|data protection commission(?:er)?|charities regulator|companies registration office)",
    re.I,
)
FOREIGN_PREFIX_RE = re.compile(r"\+\s*(?!353\b)\d{1,3}\D{0,5}$")
PRIORITY_LINK_RE = re.compile(
    r"\b(?:contact|about|location|office|support|sales|customer[\s-]*service|get[\s-]*in[\s-]*touch)\b",
    re.I,
)
LEGAL_LINK_RE = re.compile(r"\b(?:privacy|terms|legal|imprint)\b", re.I)
EXCLUDED_WEBSITE_HOSTS = {
    "linkedin.com", "facebook.com", "instagram.com", "twitter.com", "x.com",
    "youtube.com", "crunchbase.com", "pitchbook.com", "wikipedia.org",
}
GENERIC_COMPANY_WORDS = {
    "and", "company", "co", "corp", "corporation", "group", "health", "holdings",
    "inc", "ireland", "limited", "ltd", "medical", "plc", "solutions", "technology",
    "technologies", "the",
}


@dataclass(frozen=True)
class FetchResult:
    url: str
    body: str = ""
    status: int = 0
    error: str = ""
    content_type: str = ""


def fetch_url(url: str) -> FetchResult:
    request = Request(url, headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"})
    last = FetchResult(url=url)
    for attempt in range(1):
        try:
            with urlopen(request, timeout=12, context=ssl.create_default_context()) as response:
                try:
                    raw = response.read(2_000_000)
                except IncompleteRead as exc:
                    raw = exc.partial
                return FetchResult(
                    response.geturl(), raw.decode("utf-8", "ignore"),
                    getattr(response, "status", 200), "", response.headers.get("Content-Type", ""),
                )
        except HTTPError as exc:
            last = FetchResult(url, status=exc.code, error=str(exc))
            if exc.code not in {408, 425, 429, 500, 502, 503, 504}:
                return last
        except (OSError, URLError) as exc:
            last = FetchResult(url, error=str(exc))
    return last


@dataclass(frozen=True)
class CompanyInput:
    company_id: str
    company_name: str
    website: str
    geography: str
    product_type: str
    description: str
    discovery_source: str
    discovery_evidence_url: str
    source_rows: tuple[int, ...]


@dataclass(frozen=True)
class PhoneCandidate:
    phone_display: str
    phone_e164: str
    phone_type: str
    source_url: str
    extraction_method: str
    context: str
    score: int


@dataclass
class CompanyResult:
    company_id: str
    company_name: str
    geography: str
    supplied_website: str
    resolved_website: str = ""
    website_method: str = ""
    website_status: str = "missing"
    phone_status: str = "not_found"
    phone_display: str = ""
    phone_e164: str = ""
    phone_type: str = ""
    source_url: str = ""
    extraction_method: str = ""
    context: str = ""
    pages_checked: int = 0
    candidate_count: int = 0
    excluded_mobile_count: int = 0
    source_rows: str = ""
    notes: str = ""
    all_candidates: list[dict] = field(default_factory=list)
    mobile_candidates: list[dict] = field(default_factory=list)


class PageParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.text_parts: list[str] = []
        self.links: list[tuple[str, str]] = []
        self._href = ""
        self._anchor_parts: list[str] = []
        self.tel_hrefs: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "svg", "noscript"}:
            self._ignored_depth += 1
            return
        if tag != "a" or self._ignored_depth:
            return
        values = {key.lower(): value or "" for key, value in attrs}
        self._href = values.get("href", "").strip()
        self._anchor_parts = []
        if self._href.lower().startswith("tel:"):
            self.tel_hrefs.append(self._href[4:].split("?", 1)[0])

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        text = clean_space(data)
        if not text:
            return
        self.text_parts.append(text)
        if self._href:
            self._anchor_parts.append(text)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "svg", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if tag == "a" and self._href:
            self.links.append((self._href, clean_space(" ".join(self._anchor_parts))))
            self._href = ""
            self._anchor_parts = []


def clean_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def company_key(value: str) -> str:
    return clean_space(re.sub(r"[^a-z0-9]+", " ", (value or "").lower()))


def company_identity_words(value: str) -> list[str]:
    return [
        word for word in company_key(value).split()
        if len(word) > 1 and word not in GENERIC_COMPANY_WORDS
    ]


def make_company_id(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:48] or "company"
    digest = hashlib.sha1(company_key(name).encode("utf-8")).hexdigest()[:8]
    return f"ie-{slug}-{digest}"


def load_irish_companies(workbook_path: Path) -> list[CompanyInput]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Lead Intake"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    grouped: dict[str, dict] = {}
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, values))
        geography = clean_space(str(row.get("Geography") or ""))
        if "ireland" not in geography.lower():
            continue
        name = clean_space(str(row.get("Company") or ""))
        if not name:
            continue
        key = company_key(name)
        if key not in grouped:
            grouped[key] = {
                "company_id": make_company_id(name),
                "company_name": name,
                "website": clean_space(str(row.get("Website") or "")),
                "geography": geography,
                "product_type": clean_space(str(row.get("Product type") or "")),
                "description": clean_space(str(row.get("Company description") or "")),
                "discovery_source": clean_space(str(row.get("Discovery source") or "")),
                "discovery_evidence_url": clean_space(str(row.get("Discovery evidence URL") or "")),
                "source_rows": [row_number],
            }
        else:
            current = grouped[key]
            current["source_rows"].append(row_number)
            for field_name, source_name in (
                ("website", "Website"),
                ("product_type", "Product type"),
                ("description", "Company description"),
                ("discovery_evidence_url", "Discovery evidence URL"),
            ):
                if not current[field_name] and row.get(source_name):
                    current[field_name] = clean_space(str(row[source_name]))
    workbook.close()
    return [
        CompanyInput(**{**row, "source_rows": tuple(row["source_rows"])})
        for row in sorted(grouped.values(), key=lambda item: item["company_name"].lower())
    ]


def canonical_root(url: str) -> str:
    value = clean_space(url)
    if not value:
        return ""
    if not value.startswith(("http://", "https://")):
        value = "https://" + value
    parts = urlsplit(value)
    host = parts.netloc.lower().split("@")[-1].split(":", 1)[0]
    if not host:
        return ""
    return urlunsplit((parts.scheme or "https", host, "", "", ""))


def host_key(url: str) -> str:
    return urlsplit(canonical_root(url)).netloc.removeprefix("www.")


def host_is_excluded(url: str) -> bool:
    host = host_key(url)
    return not host or any(host == blocked or host.endswith("." + blocked) for blocked in EXCLUDED_WEBSITE_HOSTS)


def parse_page(html: str) -> PageParser:
    parser = PageParser()
    try:
        parser.feed(html or "")
        parser.close()
    except Exception:
        pass
    return parser


def normalize_irish_phone(raw: str) -> tuple[str, str, str] | None:
    compact = re.sub(r"[^\d+]", "", raw or "")
    if compact.startswith("00353"):
        compact = "+353" + compact[5:]
    if compact.startswith("+3530"):
        compact = "+353" + compact[5:]
    if compact.startswith("+353"):
        national = compact[4:]
    elif compact.startswith("0"):
        national = compact[1:]
    else:
        return None
    if not national.isdigit() or len(national) not in {8, 9}:
        return None
    if national.startswith("8") and not national.startswith("818"):
        phone_type = "mobile"
    elif national.startswith(("1", "2", "4", "5", "6", "7", "9", "818")):
        phone_type = "business_landline"
    else:
        return None
    e164 = "+353" + national
    display = "0" + national
    return display, e164, phone_type


def _context(text: str, start: int, end: int, radius: int = 90) -> str:
    return clean_space(text[max(0, start - radius): min(len(text), end + radius)])


def extract_phone_candidates(html: str, page_url: str) -> tuple[list[PhoneCandidate], list[PhoneCandidate], int]:
    parser = parse_page(html)
    candidates: dict[str, PhoneCandidate] = {}
    mobile_candidates: dict[str, PhoneCandidate] = {}
    mobile_occurrences = 0
    text = clean_space(" ".join(parser.text_parts))
    contexts_by_e164: dict[str, str] = {}
    for match in PHONE_RE.finditer(text):
        normalized = normalize_irish_phone(match.group(0))
        if not normalized:
            continue
        display, e164, phone_type = normalized
        context = _context(text, match.start(), match.end())
        immediate_prefix = text[max(0, match.start() - 24):match.start()]
        if FOREIGN_PREFIX_RE.search(immediate_prefix):
            continue
        if phone_type == "mobile":
            mobile_occurrences += 1
            if THIRD_PARTY_CONTEXT_RE.search(context) or FAX_CONTEXT_RE.search(immediate_prefix):
                continue
            if not PHONE_CONTEXT_RE.search(context):
                continue
            candidate = PhoneCandidate(
                display, e164, phone_type, page_url, "visible_text", context[:240],
                65 + (15 if PRIORITY_LINK_RE.search(urlsplit(page_url).path) else 0),
            )
            existing = mobile_candidates.get(e164)
            if not existing or candidate.score > existing.score:
                mobile_candidates[e164] = candidate
            continue
        if FAX_CONTEXT_RE.search(immediate_prefix):
            continue
        if THIRD_PARTY_CONTEXT_RE.search(context):
            continue
        if not PHONE_CONTEXT_RE.search(context):
            continue
        contexts_by_e164[e164] = context
        page_bonus = 15 if PRIORITY_LINK_RE.search(urlsplit(page_url).path) else 0
        candidate = PhoneCandidate(
            display, e164, phone_type, page_url, "visible_text", context[:240], 60 + page_bonus,
        )
        existing = candidates.get(e164)
        if not existing or candidate.score > existing.score:
            candidates[e164] = candidate
    for raw in parser.tel_hrefs:
        normalized = normalize_irish_phone(raw)
        if not normalized:
            continue
        display, e164, phone_type = normalized
        if phone_type == "mobile":
            mobile_occurrences += 1
            context = contexts_by_e164.get(e164, "Telephone link")
            if THIRD_PARTY_CONTEXT_RE.search(context):
                continue
            mobile_candidates[e164] = PhoneCandidate(
                display, e164, phone_type, page_url, "tel_link", context[:240], 100,
            )
            continue
        context = contexts_by_e164.get(e164, "Telephone link")
        if THIRD_PARTY_CONTEXT_RE.search(context):
            continue
        candidates[e164] = PhoneCandidate(
            display, e164, phone_type, page_url, "tel_link", context[:240], 100,
        )
    return (
        sorted(candidates.values(), key=lambda row: (-row.score, row.phone_e164)),
        sorted(mobile_candidates.values(), key=lambda row: (-row.score, row.phone_e164)),
        mobile_occurrences,
    )


def prioritized_internal_links(html: str, page_url: str, limit: int = 5) -> list[str]:
    base_host = host_key(page_url)
    scored: dict[str, int] = {}
    for href, anchor in parse_page(html).links:
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        absolute = urljoin(page_url, href).split("#", 1)[0]
        if host_key(absolute) != base_host:
            continue
        parts = urlsplit(absolute)
        if parts.path.lower().endswith((".pdf", ".jpg", ".jpeg", ".png", ".zip")):
            continue
        text = f"{parts.path} {anchor}"
        score = 0
        if PRIORITY_LINK_RE.search(text):
            score = 20
        elif LEGAL_LINK_RE.search(text):
            score = 5
        if score:
            scored[absolute] = max(scored.get(absolute, 0), score)
    return [url for url, _ in sorted(scored.items(), key=lambda item: (-item[1], item[0]))[:limit]]


def page_matches_company(company_name: str, html: str, url: str) -> bool:
    words = company_identity_words(company_name)
    if not words:
        return False
    text = company_key(" ".join(parse_page(html[:250000]).text_parts)[:30000])
    domain = re.sub(r"[^a-z0-9]", "", host_key(url).split(".", 1)[0])
    if len(words) == 1:
        word = words[0]
        return word in text.split() and (word in domain or company_key(company_name) in text)
    matches = sum(word in text.split() for word in words)
    return matches >= max(1, min(2, len(words)))


def candidate_website_score(company_name: str, url: str, anchor: str = "", title: str = "", snippet: str = "") -> int:
    if host_is_excluded(url):
        return -100
    words = company_identity_words(company_name)
    if not words:
        return 0
    haystack = company_key(f"{anchor} {title} {snippet}")
    domain = re.sub(r"[^a-z0-9]", "", host_key(url).split(".", 1)[0])
    score = sum(2 for word in words if word in haystack.split())
    compact_name = "".join(words)
    if compact_name and compact_name in domain:
        score += 6
    elif any(len(word) >= 5 and word in domain for word in words):
        score += 3
    if "official" in haystack:
        score += 1
    return score


def probable_domains(company_name: str) -> list[str]:
    words = company_identity_words(company_name)
    if not words:
        return []
    labels = ["".join(words)]
    if len(words) > 1:
        labels.append("-".join(words))
    return [f"https://{label}.{tld}" for label in labels for tld in ("ie", "com")]


class CachedFetcher:
    def __init__(self, delegate: Callable[[str], FetchResult] = fetch_url) -> None:
        self.delegate = delegate
        self._cache: dict[str, FetchResult] = {}
        self._lock = threading.Lock()

    def __call__(self, url: str) -> FetchResult:
        key = url.split("#", 1)[0]
        with self._lock:
            cached = self._cache.get(key)
        if cached is not None:
            return cached
        result = self.delegate(key)
        with self._lock:
            self._cache[key] = result
        return result


class BoundedSearch:
    """Serialize free search and cap it at one attempt per company."""

    def __init__(
        self,
        delegate: Callable[[str], tuple[list[PublicSearchHit], str | None]] | None = None,
        error_limit: int = 2,
    ) -> None:
        self.delegate = delegate
        self.error_limit = error_limit
        self._consecutive_errors = 0
        self._disabled_error = ""
        self._lock = threading.Lock()

    def __call__(self, query: str) -> tuple[list[PublicSearchHit], str | None]:
        with self._lock:
            if self._disabled_error:
                return [], self._disabled_error
            if self.delegate is not None:
                hits, error = self.delegate(query)
            else:
                hits, error = duckduckgo_search(query, attempts=1)
            if hits:
                self._consecutive_errors = 0
            elif error:
                self._consecutive_errors += 1
                if self._consecutive_errors >= self.error_limit:
                    self._disabled_error = f"Search disabled after {self.error_limit} consecutive errors: {error}"
                    return [], self._disabled_error
            return hits, error


def _verified_root(company: CompanyInput, url: str, fetcher: Callable[[str], FetchResult]) -> str:
    root = canonical_root(url)
    if not root or host_is_excluded(root):
        return ""
    fetched = fetcher(root)
    if fetched.error or fetched.status >= 400 or not fetched.body:
        return ""
    final_root = canonical_root(fetched.url or root)
    return final_root if page_matches_company(company.company_name, fetched.body, final_root) else ""


def recover_website(
    company: CompanyInput,
    fetcher: Callable[[str], FetchResult],
    search_fn: Callable[[str], tuple[list[PublicSearchHit], str | None]] = configured_search,
) -> tuple[str, str, str]:
    if company.website:
        root = canonical_root(company.website)
        return root, "supplied_website", "supplied" if root else "invalid"

    if company.discovery_evidence_url:
        evidence = fetcher(company.discovery_evidence_url)
        if not evidence.error and evidence.status < 400 and evidence.body:
            links = parse_page(evidence.body).links
            ranked = sorted(
                (
                    (candidate_website_score(company.company_name, urljoin(evidence.url or company.discovery_evidence_url, href), anchor),
                     urljoin(evidence.url or company.discovery_evidence_url, href))
                    for href, anchor in links
                ),
                key=lambda item: (-item[0], item[1]),
            )
            for score, url in ranked[:8]:
                if score < 3 or host_key(url) == host_key(company.discovery_evidence_url):
                    continue
                root = _verified_root(company, url, fetcher)
                if root:
                    return root, "discovery_evidence_link", "resolved"

    for candidate in probable_domains(company.company_name):
        root = _verified_root(company, candidate, fetcher)
        if root:
            return root, "direct_domain_probe", "resolved"

    query = f'"{company.company_name}" Ireland official website'
    hits, error = search_fn(query)
    for hit in sorted(
        hits[:10],
        key=lambda row: -candidate_website_score(company.company_name, row.url, title=row.title, snippet=row.snippet),
    ):
        if candidate_website_score(company.company_name, hit.url, title=hit.title, snippet=hit.snippet) < 4:
            continue
        root = _verified_root(company, hit.url, fetcher)
        if root:
            return root, "web_search", "resolved"
    return "", "web_search", "search_error" if error and not hits else "not_found"


def crawl_company(
    company: CompanyInput,
    fetcher: Callable[[str], FetchResult],
    search_fn: Callable[[str], tuple[list[PublicSearchHit], str | None]] = configured_search,
    max_pages: int = 6,
) -> CompanyResult:
    result = CompanyResult(
        company_id=company.company_id,
        company_name=company.company_name,
        geography=company.geography,
        supplied_website=company.website,
        source_rows=", ".join(str(row) for row in company.source_rows),
    )
    website, method, website_status = recover_website(company, fetcher, search_fn)
    result.resolved_website = website
    result.website_method = method
    result.website_status = website_status
    if not website:
        result.notes = "No verified official website was available for phone extraction"
        return result

    queue = [website]
    visited: set[str] = set()
    candidates: dict[str, PhoneCandidate] = {}
    mobile_candidates: dict[str, PhoneCandidate] = {}
    errors: list[str] = []
    while queue and len(visited) < max_pages:
        page_url = queue.pop(0)
        if page_url in visited:
            continue
        visited.add(page_url)
        fetched = fetcher(page_url)
        if fetched.error or fetched.status >= 400 or not fetched.body:
            errors.append(f"{page_url}: {fetched.error or 'HTTP ' + str(fetched.status)}")
            continue
        actual_url = fetched.url or page_url
        found, found_mobiles, mobile_occurrences = extract_phone_candidates(fetched.body, actual_url)
        result.excluded_mobile_count += mobile_occurrences
        for candidate in found:
            existing = candidates.get(candidate.phone_e164)
            if not existing or candidate.score > existing.score:
                candidates[candidate.phone_e164] = candidate
        for candidate in found_mobiles:
            existing = mobile_candidates.get(candidate.phone_e164)
            if not existing or candidate.score > existing.score:
                mobile_candidates[candidate.phone_e164] = candidate
        if len(visited) == 1:
            queue.extend(link for link in prioritized_internal_links(fetched.body, actual_url, max_pages - 1) if link not in visited)
    result.pages_checked = len(visited)
    ranked = sorted(candidates.values(), key=lambda row: (-row.score, row.phone_e164))
    result.candidate_count = len(ranked)
    result.all_candidates = [asdict(row) for row in ranked]
    result.mobile_candidates = [
        asdict(row) for row in sorted(mobile_candidates.values(), key=lambda row: (-row.score, row.phone_e164))
    ]
    if ranked:
        best = ranked[0]
        result.phone_status = "found"
        result.phone_display = best.phone_display
        result.phone_e164 = best.phone_e164
        result.phone_type = best.phone_type
        result.source_url = best.source_url
        result.extraction_method = best.extraction_method
        result.context = best.context
    elif errors and len(errors) == len(visited):
        result.phone_status = "site_error"
        result.notes = errors[0][:500]
    else:
        result.notes = "No public Irish non-mobile phone number found on checked official pages"
    return result


def write_results(output_dir: Path, companies: list[CompanyInput], results: list[CompanyResult], run_date: str) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    result_rows = [asdict(row) for row in results]
    mobile_rows = [
        {
            "company_id": row.company_id,
            "company_name": row.company_name,
            "geography": row.geography,
            "resolved_website": row.resolved_website,
            "source_rows": row.source_rows,
            **candidate,
        }
        for row in results for candidate in row.mobile_candidates
    ]
    summary = {
        "schema_version": "1.0",
        "generated_at": run_date,
        "irish_source_rows": sum(len(row.source_rows) for row in companies),
        "distinct_companies": len(companies),
        "companies_with_supplied_website": sum(bool(row.website) for row in companies),
        "website_status_counts": dict(Counter(row.website_status for row in results)),
        "phone_status_counts": dict(Counter(row.phone_status for row in results)),
        "phones_found": sum(row.phone_status == "found" for row in results),
        "mobile_occurrences_found": sum(row.excluded_mobile_count for row in results),
        "mobile_company_number_rows": len(mobile_rows),
        "unique_mobile_numbers": len({row["phone_e164"] for row in mobile_rows}),
        "companies_with_mobile": len({row["company_id"] for row in mobile_rows}),
    }
    files = {
        "results_json": output_dir / "irish_phone_results.json",
        "results_csv": output_dir / "irish_phone_results.csv",
        "mobiles_csv": output_dir / "irish_mobile_results.csv",
        "summary": output_dir / "run_summary.json",
    }
    files["results_json"].write_text(
        json.dumps({"summary": summary, "records": result_rows, "mobile_records": mobile_rows}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    files["summary"].write_text(json.dumps(summary, indent=2), encoding="utf-8")
    csv_fields = [name for name in CompanyResult.__dataclass_fields__ if name != "all_candidates"]
    with files["results_csv"].open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=csv_fields)
        writer.writeheader()
        for row in result_rows:
            writer.writerow({name: row[name] for name in csv_fields})
    mobile_fields = [
        "company_id", "company_name", "geography", "phone_display", "phone_e164", "phone_type",
        "source_url", "extraction_method", "context", "score", "resolved_website", "source_rows",
    ]
    with files["mobiles_csv"].open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=mobile_fields)
        writer.writeheader()
        writer.writerows(mobile_rows)
    return files


def run_extraction(
    workbook_path: Path,
    output_dir: Path,
    workers: int = 8,
    max_pages: int = 6,
    limit: int | None = None,
    fetcher: Callable[[str], FetchResult] = fetch_url,
    search_fn: Callable[[str], tuple[list[PublicSearchHit], str | None]] | None = None,
    run_date: str | None = None,
    progress_fn: Callable[[int, int, CompanyResult], None] | None = None,
) -> tuple[dict, dict[str, Path]]:
    run_date = run_date or date.today().isoformat()
    companies = load_irish_companies(workbook_path)
    if limit is not None:
        companies = companies[:limit]
    cached_fetcher = CachedFetcher(fetcher)
    bounded_search = BoundedSearch(search_fn)
    results: list[CompanyResult] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {
            pool.submit(crawl_company, company, cached_fetcher, bounded_search, max_pages): company
            for company in companies
        }
        for future in as_completed(futures):
            company = futures[future]
            try:
                results.append(future.result())
            except Exception as exc:
                results.append(CompanyResult(
                    company_id=company.company_id,
                    company_name=company.company_name,
                    geography=company.geography,
                    supplied_website=company.website,
                    source_rows=", ".join(str(row) for row in company.source_rows),
                    phone_status="extractor_error",
                    notes=f"{type(exc).__name__}: {exc}"[:500],
                ))
            if progress_fn:
                progress_fn(len(results), len(companies), results[-1])
    results.sort(key=lambda row: row.company_name.lower())
    files = write_results(output_dir, companies, results, run_date)
    summary = json.loads(files["summary"].read_text(encoding="utf-8"))
    return summary, files
